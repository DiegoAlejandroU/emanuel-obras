"""Reportes exportables (PDF/XLSX) de indicadores de avance para gerencia e
interventoría.

Se construyen a partir de los mismos datos que ya expone `/indicadores` y
`/alertas` — este módulo solo los da en formato de archivo. No agrega
reglas de negocio nuevas ni escribe en la base de datos.
"""

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from .. import models
from .alertas_service import listar_alertas
from .obras_service import calcular_indicadores, obtener_obra_o_404

ETIQUETAS_TIPO_ALERTA = {
    "retraso_fisico": "Retraso físico",
    "retraso_financiero": "Retraso financiero",
    "incidente_grave": "Incidente grave",
}

_RELLENO_ENCABEZADO = PatternFill("solid", fgColor="4B5563")


def _nombre_archivo_base(obra: models.Obra) -> str:
    slug = "".join(c if c.isalnum() else "_" for c in obra.nombre.strip().lower())
    fecha = datetime.now(timezone.utc).strftime("%Y%m%d")
    return f"reporte_obra_{obra.id}_{slug}_{fecha}"


def _datos_reporte(db: Session, obra_id: int):
    obra = obtener_obra_o_404(db, obra_id)
    indicadores = calcular_indicadores(db, obra_id)
    alertas = listar_alertas(db, obra_id)
    return obra, indicadores, alertas


# ---------- PDF ----------


def generar_reporte_pdf(db: Session, obra_id: int) -> tuple[bytes, str]:
    obra, indicadores, alertas = _datos_reporte(db, obra_id)

    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        title=f"Reporte de avance — {obra.nombre}",
    )

    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle("TituloReporte", parent=estilos["Title"], fontSize=16, spaceAfter=4)
    estilo_subtitulo = ParagraphStyle("Subtitulo", parent=estilos["Normal"], textColor=colors.grey, spaceAfter=16)
    estilo_seccion = ParagraphStyle("Seccion", parent=estilos["Heading2"], spaceBefore=16, spaceAfter=8)

    elementos = [
        Paragraph(f"Reporte de avance — {obra.nombre}", estilo_titulo),
        Paragraph(
            f"Generado el {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC · "
            f"Emanuel Ingeniería y Construcciones S.A.S.",
            estilo_subtitulo,
        ),
    ]

    datos_generales = [
        ["Contratista", obra.contratista or "—"],
        ["Ubicación", obra.ubicacion or "—"],
        ["Estado", obra.estado],
        ["Fecha de inicio", str(obra.fecha_inicio) if obra.fecha_inicio else "—"],
        ["Fecha fin estimada", str(obra.fecha_fin_estimada) if obra.fecha_fin_estimada else "—"],
        ["Presupuesto total", f"${indicadores.presupuesto_total:,.0f}"],
    ]
    tabla_general = Table(datos_generales, colWidths=[5 * cm, 10 * cm])
    tabla_general.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#e5e7eb")),
            ]
        )
    )
    elementos.append(tabla_general)

    elementos.append(Paragraph("Indicadores de avance", estilo_seccion))
    tabla_indicadores = Table(
        [
            ["Avance físico", "Avance financiero"],
            [f"{indicadores.avance_fisico_porcentual}%", f"{indicadores.avance_financiero_porcentual}%"],
        ],
        colWidths=[7.5 * cm, 7.5 * cm],
    )
    tabla_indicadores.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 1), (-1, 1), 14),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#15803d")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e5e7eb")),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elementos.append(tabla_indicadores)

    elementos.append(Paragraph("Avance por actividad", estilo_seccion))
    filas_actividades = [["Actividad", "Peso", "Avance acumulado"]]
    for act in indicadores.actividades:
        filas_actividades.append(
            [act.nombre, f"{round(act.peso_porcentual * 100)}%", f"{act.avance_acumulado_porcentual}%"]
        )
    if len(filas_actividades) == 1:
        filas_actividades.append(["Sin actividades registradas", "—", "—"])
    tabla_actividades = Table(filas_actividades, colWidths=[9 * cm, 3 * cm, 3 * cm])
    tabla_actividades.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e5e7eb")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elementos.append(tabla_actividades)

    elementos.append(Paragraph("Alertas", estilo_seccion))
    filas_alertas = [["Tipo", "Mensaje", "Fecha", "Estado"]]
    for alerta in alertas:
        filas_alertas.append(
            [
                ETIQUETAS_TIPO_ALERTA.get(alerta.tipo, alerta.tipo),
                Paragraph(alerta.mensaje, estilos["Normal"]),
                alerta.fecha_generacion.strftime("%Y-%m-%d") if alerta.fecha_generacion else "—",
                "Activa" if alerta.estado == "activa" else "Resuelta",
            ]
        )
    if len(filas_alertas) == 1:
        filas_alertas.append(["—", "Sin alertas registradas para esta obra.", "—", "—"])
    tabla_alertas = Table(filas_alertas, colWidths=[3 * cm, 7 * cm, 2.5 * cm, 2.5 * cm])
    tabla_alertas.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e5e7eb")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elementos.append(tabla_alertas)
    elementos.append(Spacer(1, 12))

    documento.build(elementos)
    return buffer.getvalue(), f"{_nombre_archivo_base(obra)}.pdf"


# ---------- XLSX ----------


def _autoajustar_columnas(hoja, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def generar_reporte_xlsx(db: Session, obra_id: int) -> tuple[bytes, str]:
    obra, indicadores, alertas = _datos_reporte(db, obra_id)

    libro = Workbook()
    fuente_encabezado = Font(bold=True, color="FFFFFF")
    fuente_etiqueta = Font(bold=True)

    # ---- Hoja Resumen ----
    resumen = libro.active
    resumen.title = "Resumen"
    resumen["A1"] = f"Reporte de avance — {obra.nombre}"
    resumen["A1"].font = Font(bold=True, size=14)
    resumen["A2"] = f"Generado el {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC"
    resumen["A2"].font = Font(italic=True, color="6B7280")

    filas_resumen = [
        ("Contratista", obra.contratista or "—"),
        ("Ubicación", obra.ubicacion or "—"),
        ("Estado", obra.estado),
        ("Fecha de inicio", str(obra.fecha_inicio) if obra.fecha_inicio else "—"),
        ("Fecha fin estimada", str(obra.fecha_fin_estimada) if obra.fecha_fin_estimada else "—"),
        ("Presupuesto total", indicadores.presupuesto_total),
        ("Avance físico (%)", indicadores.avance_fisico_porcentual),
        ("Avance financiero (%)", indicadores.avance_financiero_porcentual),
    ]
    fila_actual = 4
    for etiqueta, valor in filas_resumen:
        resumen.cell(row=fila_actual, column=1, value=etiqueta).font = fuente_etiqueta
        resumen.cell(row=fila_actual, column=2, value=valor)
        fila_actual += 1
    _autoajustar_columnas(resumen, [22, 28])

    # ---- Hoja Actividades ----
    hoja_actividades = libro.create_sheet("Actividades")
    encabezados = ["Actividad", "Peso (%)", "Avance acumulado (%)"]
    for col, texto in enumerate(encabezados, start=1):
        celda = hoja_actividades.cell(row=1, column=col, value=texto)
        celda.font = fuente_encabezado
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center")
    for fila, act in enumerate(indicadores.actividades, start=2):
        hoja_actividades.cell(row=fila, column=1, value=act.nombre)
        hoja_actividades.cell(row=fila, column=2, value=round(act.peso_porcentual * 100, 2))
        hoja_actividades.cell(row=fila, column=3, value=act.avance_acumulado_porcentual)
    _autoajustar_columnas(hoja_actividades, [40, 12, 20])

    # ---- Hoja Alertas ----
    hoja_alertas = libro.create_sheet("Alertas")
    encabezados_alertas = ["Tipo", "Mensaje", "Fecha", "Estado"]
    for col, texto in enumerate(encabezados_alertas, start=1):
        celda = hoja_alertas.cell(row=1, column=col, value=texto)
        celda.font = fuente_encabezado
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center")
    for fila, alerta in enumerate(alertas, start=2):
        hoja_alertas.cell(row=fila, column=1, value=ETIQUETAS_TIPO_ALERTA.get(alerta.tipo, alerta.tipo))
        hoja_alertas.cell(row=fila, column=2, value=alerta.mensaje)
        hoja_alertas.cell(
            row=fila, column=3, value=alerta.fecha_generacion.strftime("%Y-%m-%d") if alerta.fecha_generacion else "—"
        )
        hoja_alertas.cell(row=fila, column=4, value="Activa" if alerta.estado == "activa" else "Resuelta")
    _autoajustar_columnas(hoja_alertas, [20, 55, 14, 14])

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue(), f"{_nombre_archivo_base(obra)}.xlsx"
